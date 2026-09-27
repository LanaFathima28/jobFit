import io
import pytest
import docx
import fitz  # PyMuPDF
from fastapi.testclient import TestClient

from app.main import app
from app.services.extraction_service import (
    extract_text_from_file_bytes,
    ExtractionStatus,
)
from app.services.text_cleaner import clean_extracted_text

client = TestClient(app)


# --- Helper Functions to Generate Sample In-Memory Files ---

def create_sample_pdf(text: str) -> bytes:
    """Generates a valid PDF document with text using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(paragraphs: list[str], table_data: list[list[str]] = None) -> bytes:
    """Generates a valid DOCX document using python-docx."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --- Unit Tests: Text Cleaner ---

def test_text_cleaner():
    raw = "  John Doe  \n\n\n\nPage 1 of 3\nSenior  Developer\xa0\xa0\n\n\nWork History  "
    cleaned = clean_extracted_text(raw)

    assert "Page 1 of 3" not in cleaned
    assert "John Doe" in cleaned
    assert "Senior Developer" in cleaned
    assert "\n\n\n" not in cleaned


# --- Unit Tests: Extraction Service ---

def test_extract_txt_file():
    txt_bytes = b"Software Engineer Resume\nSkills: Python, FastAPI, Docker."
    result = extract_text_from_file_bytes(txt_bytes, "resume.txt")

    assert result.status == ExtractionStatus.SUCCESS
    assert "Software Engineer Resume" in result.text
    assert "Python, FastAPI" in result.text


def test_extract_pdf_file():
    pdf_bytes = create_sample_pdf("Jane Smith - Senior Software Architect\n5+ years of experience in Python.")
    result = extract_text_from_file_bytes(pdf_bytes, "jane_smith.pdf")

    assert result.status == ExtractionStatus.SUCCESS
    assert "Jane Smith" in result.text
    assert "Software Architect" in result.text


def test_extract_docx_file():
    docx_bytes = create_sample_docx(
        paragraphs=["Alex Johnson", "Backend Engineer"],
        table_data=[["Company", "Role"], ["Acme Corp", "Tech Lead"]]
    )
    result = extract_text_from_file_bytes(docx_bytes, "alex_johnson.docx")

    assert result.status == ExtractionStatus.SUCCESS
    assert "Alex Johnson" in result.text
    assert "Acme Corp" in result.text
    assert "Tech Lead" in result.text


def test_extract_corrupt_pdf():
    corrupt_bytes = b"%PDF-1.4 Corrupted content invalid header !!! \x00\xff\xfe"
    result = extract_text_from_file_bytes(corrupt_bytes, "corrupted_resume.pdf")

    assert result.status == ExtractionStatus.FAILED
    assert result.error_message is not None
    assert "No extractable text" in result.error_message or "failed" in result.error_message.lower()


def test_extract_empty_file():
    empty_bytes = b""
    result = extract_text_from_file_bytes(empty_bytes, "empty_file.pdf")

    assert result.status == ExtractionStatus.FAILED
    assert "empty" in result.error_message.lower()


def test_extract_unsupported_format():
    exe_bytes = b"MZ executable header binary data"
    result = extract_text_from_file_bytes(exe_bytes, "malware.exe")

    assert result.status == ExtractionStatus.FAILED
    assert "Unsupported file extension" in result.error_message


# --- API Endpoint Integration Tests ---

def test_candidate_upload_and_retrieval_api():
    pdf_bytes = create_sample_pdf("Candidate: Robert Vance\nSpecialty: PostgreSQL, Python, Vector Search.")
    
    # 1. Single Resume Upload
    response = client.post(
        "/api/v1/candidates/upload",
        files={"file": ("robert_vance.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 201
    data = response.json()
    assert isinstance(data, list)
    assert len(data) == 1
    candidate_id = data[0]["id"]
    assert data[0]["original_filename"] == "robert_vance.pdf"
    assert data[0]["extraction_status"] == "success"
    assert "Robert Vance" in data[0]["raw_text"]

    # 2. Get Candidate Detail
    get_res = client.get(f"/api/v1/candidates/{candidate_id}")
    assert get_res.status_code == 200
    assert get_res.json()["id"] == candidate_id

    # 3. Bulk Upload
    pdf2 = create_sample_pdf("Candidate Two\nSkills: React, Next.js")
    docx1 = create_sample_docx(["Candidate Three", "Skills: Docker, Kubernetes"])
    
    bulk_res = client.post(
        "/api/v1/candidates/upload",
        files=[
            ("files", ("candidate_two.pdf", pdf2, "application/pdf")),
            ("files", ("candidate_three.docx", docx1, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
    )
    assert bulk_res.status_code == 201
    bulk_data = bulk_res.json()
    assert len(bulk_data) == 2

    # 4. List Candidates
    list_res = client.get("/api/v1/candidates")
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 3


def test_corrupt_candidate_upload_api():
    corrupt_bytes = b"NOT A VALID PDF OR ZIP HEADER"
    response = client.post(
        "/api/v1/candidates/upload",
        files={"file": ("corrupt_user.pdf", corrupt_bytes, "application/pdf")}
    )
    # The API stores DB record even on extraction failure
    assert response.status_code == 201
    data = response.json()
    assert data[0]["extraction_status"] == "failed"
    assert data[0]["extraction_error"] is not None


def test_job_upload_and_retrieval_api():
    job_pdf = create_sample_pdf("Senior DevOps Engineer\nRequirements: AWS, TerraForm, CI/CD pipelines.")

    # 1. File Upload
    response = client.post(
        "/api/v1/jobs/upload",
        files={"file": ("devops_engineer.pdf", job_pdf, "application/pdf")},
        data={"title": "Senior DevOps Engineer"}
    )
    assert response.status_code == 201
    job_data = response.json()
    assert job_data["title"] == "Senior DevOps Engineer"
    assert job_data["extraction_status"] == "success"
    assert "DevOps" in job_data["description_text"]
    job_id = job_data["id"]

    # 2. Get Job Detail
    get_res = client.get(f"/api/v1/jobs/{job_id}")
    assert get_res.status_code == 200
    assert get_res.json()["id"] == job_id

    # 3. Raw Text Upload
    raw_res = client.post(
        "/api/v1/jobs/upload",
        data={
            "title": "Lead Python Developer",
            "raw_text": "We are seeking a Lead Python Developer with 7+ years experience in FastAPI and microservices."
        }
    )
    assert raw_res.status_code == 201
    assert raw_res.json()["file_type"] == "raw_text"
    assert raw_res.json()["extraction_status"] == "success"

    # 4. List Jobs
    list_res = client.get("/api/v1/jobs")
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 2
