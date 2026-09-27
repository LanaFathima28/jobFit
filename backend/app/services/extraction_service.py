import io
import logging
from dataclasses import dataclass, field
from enum import Enum
from typing import Optional

import docx
import fitz  # PyMuPDF
import pdfplumber

from app.services.text_cleaner import clean_extracted_text

logger = logging.getLogger(__name__)


class ExtractionStatus(str, Enum):
    SUCCESS = "success"
    PARTIAL = "partial"
    FAILED = "failed"


@dataclass
class ExtractionResult:
    text: str = ""
    status: ExtractionStatus = ExtractionStatus.FAILED
    error_message: Optional[str] = None
    page_count: int = 0
    char_count: int = 0


def extract_pdf_with_pdfplumber(file_bytes: bytes) -> tuple[str, int]:
    """
    Extract text using pdfplumber. Returns (text, page_count).
    """
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text_chunks.append(extracted)
    return "\n".join(text_chunks), page_count


def extract_pdf_with_pymupdf(file_bytes: bytes) -> tuple[str, int]:
    """
    Fallback text extraction using PyMuPDF (fitz). Returns (text, page_count).
    """
    text_chunks = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    page_count = len(doc)
    for page in doc:
        extracted = page.get_text("text")
        if extracted:
            text_chunks.append(extracted)
    doc.close()
    return "\n".join(text_chunks), page_count


def extract_text_from_pdf(file_bytes: bytes) -> ExtractionResult:
    """
    Attempts PDF text extraction with pdfplumber, falling back to PyMuPDF if pdfplumber fails or produces empty text.
    """
    raw_text = ""
    page_count = 0
    primary_failed = False
    primary_error = None

    # Try pdfplumber first
    try:
        raw_text, page_count = extract_pdf_with_pdfplumber(file_bytes)
    except Exception as e:
        logger.warning(f"pdfplumber extraction failed: {str(e)}. Attempting PyMuPDF fallback.")
        primary_failed = True
        primary_error = str(e)

    # Fallback to PyMuPDF if pdfplumber failed or produced empty text
    if primary_failed or not raw_text or len(raw_text.strip()) < 10:
        try:
            fallback_text, fallback_pages = extract_pdf_with_pymupdf(file_bytes)
            if fallback_text and len(fallback_text.strip()) >= 10:
                raw_text = fallback_text
                page_count = max(page_count, fallback_pages)
            elif not raw_text and fallback_text:
                raw_text = fallback_text
                page_count = max(page_count, fallback_pages)
        except Exception as fallback_err:
            logger.warning(f"PyMuPDF fallback extraction failed: {str(fallback_err)}")
            if primary_failed:
                return ExtractionResult(
                    text="",
                    status=ExtractionStatus.FAILED,
                    error_message=f"PDF extraction failed. File may be corrupted or password-protected ({primary_error})"
                )

    cleaned_text = clean_extracted_text(raw_text)

    if not cleaned_text or len(cleaned_text) < 10:
        return ExtractionResult(
            text="",
            status=ExtractionStatus.FAILED,
            error_message="No extractable text found in PDF (scanned image-only PDF without OCR text layer)",
            page_count=page_count,
            char_count=0
        )

    return ExtractionResult(
        text=cleaned_text,
        status=ExtractionStatus.SUCCESS,
        page_count=page_count,
        char_count=len(cleaned_text)
    )


def extract_text_from_docx(file_bytes: bytes) -> ExtractionResult:
    """
    Extract text from DOCX file bytes including paragraphs and table contents.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_chunks = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_chunks.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))

        raw_text = "\n".join(text_chunks)
        cleaned_text = clean_extracted_text(raw_text)

        if not cleaned_text:
            return ExtractionResult(
                text="",
                status=ExtractionStatus.FAILED,
                error_message="DOCX document contains no readable text content.",
                char_count=0
            )

        return ExtractionResult(
            text=cleaned_text,
            status=ExtractionStatus.SUCCESS,
            page_count=1,
            char_count=len(cleaned_text)
        )
    except Exception as e:
        logger.error(f"DOCX extraction error: {str(e)}")
        return ExtractionResult(
            text="",
            status=ExtractionStatus.FAILED,
            error_message=f"Failed to read DOCX file: File may be corrupted or invalid format ({str(e)})"
        )


def extract_text_from_txt(file_bytes: bytes) -> ExtractionResult:
    """
    Extract text from plain TXT file bytes.
    """
    try:
        try:
            raw_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raw_text = file_bytes.decode("latin-1", errors="ignore")

        cleaned_text = clean_extracted_text(raw_text)
        if not cleaned_text:
            return ExtractionResult(
                text="",
                status=ExtractionStatus.FAILED,
                error_message="Text file is empty.",
                char_count=0
            )

        return ExtractionResult(
            text=cleaned_text,
            status=ExtractionStatus.SUCCESS,
            page_count=1,
            char_count=len(cleaned_text)
        )
    except Exception as e:
        return ExtractionResult(
            text="",
            status=ExtractionStatus.FAILED,
            error_message=f"Failed to decode TXT file: {str(e)}"
        )


def extract_text_from_file_bytes(file_bytes: bytes, filename: str) -> ExtractionResult:
    """
    Main extraction service entrypoint.
    Determines extraction strategy by file extension, executes parsing, cleans text,
    and returns a structured ExtractionResult.
    """
    if not file_bytes or len(file_bytes) == 0:
        return ExtractionResult(
            text="",
            status=ExtractionStatus.FAILED,
            error_message="Uploaded file is empty (0 bytes)."
        )

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        return extract_text_from_docx(file_bytes)
    elif ext in ["txt", "text", "md"]:
        return extract_text_from_txt(file_bytes)
    else:
        return ExtractionResult(
            text="",
            status=ExtractionStatus.FAILED,
            error_message=f"Unsupported file extension '.{ext}'. Supported formats: PDF, DOCX, TXT."
        )
